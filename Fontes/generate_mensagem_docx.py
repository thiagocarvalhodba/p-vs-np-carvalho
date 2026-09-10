import docx
from docx.shared import Pt, RGBColor

def create_docx():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading(level=1)
    run_title = title.add_run("Carta de Encaminhamento ao Parecer nº 09")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)
    
    with open(r"C:\MathDoCarvalho\MensagemParaOAvaliador.txt", "r", encoding="utf-8") as f:
        text = f.read()
    
    paragraphs = text.split("\n\n")
    for para in paragraphs:
        p = doc.add_paragraph()
        p_text = para.strip()
        if p_text.startswith("--------------------------------------------------------------------------------"):
            continue
        if p_text.startswith("Assunto:") or p_text.startswith("Destinatário:") or p_text.startswith("ARQUIVOS ANEXOS:"):
            run = p.add_run(p_text)
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.name = "Calibri"
        elif p_text.startswith("1. ") or p_text.startswith("2. ") or p_text.startswith("3. ") or p_text.startswith("4. ") or p_text.startswith("5. ") or p_text.startswith("6. "):
            lines = p_text.split("\n")
            run_head = p.add_run(lines[0] + "\n")
            run_head.font.bold = True
            run_head.font.size = Pt(11)
            run_head.font.name = "Calibri"
            if len(lines) > 1:
                run_body = p.add_run("\n".join(lines[1:]))
                run_body.font.size = Pt(11)
                run_body.font.name = "Calibri"
        else:
            run = p.add_run(p_text)
            run.font.size = Pt(11)
            run.font.name = "Calibri"
            
    doc.save(r"C:\MathDoCarvalho\MensagemParaOAvaliador.docx")
    print("DOCX gerado com sucesso em C:\\MathDoCarvalho\\MensagemParaOAvaliador.docx")

if __name__ == "__main__":
    create_docx()

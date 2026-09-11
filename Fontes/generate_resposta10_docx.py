import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

def clean_math(text):
    """Converts LaTeX math markup to clean Unicode math notation for Word."""
    # Replace math expressions and symbols
    repl = [
        (r'\\mathcal\{U\}_N(\([^)]+\))?', r'U_N\1'),
        (r'\\mathcal\{U\}_N', 'U_N'),
        (r'\\mathcal\{X\}', 'X'),
        (r'\\mathcal\{F\}', 'F'),
        (r'\\mathcal\{O\}_s', 'O_s'),
        (r'\\mathcal\{O\}', 'O'),
        (r'\\mathcal\{D\}_\{?\\text\{proj\}\}?', 'D_proj'),
        (r'\\mathcal\{D\}', 'D'),
        (r'\\mathcal\{C\}_0', 'C_0'),
        (r'\\mathcal\{C\}_\{?\\text\{spur\}\}?', 'C_spur'),
        (r'\\mathcal\{S\}_\{?\\text\{spur\}\}?', 'S_spur'),
        (r'\\mathcal\{B\}_\{?\\text\{spur\}\}?', 'B_spur'),
        (r'\\mathcal\{M\}_\{?\\text\{spur\}\}?', 'M_spur'),
        (r'\\Phi_\{?\\text\{quad\}\}?', 'Φ_quad'),
        (r'\\Phi_\{?\\text\{mult\}\}?', 'Φ_mult'),
        (r'\\Phi_\{?\\text\{soft\}\}?', 'Φ_soft'),
        (r'\\Phi', 'Φ'),
        (r'\\nabla\^2', '∇²'),
        (r'\\nabla', '∇'),
        (r'\\Delta_\{?\\mathcal\{F\}\}?', 'Δ_F'),
        (r'\\Delta', 'Δ'),
        (r'\\mathbf\{0\}', '0'),
        (r'\\mathbf\{1\}', '1'),
        (r'\\equiv', ' ≡ '),
        (r'\\subset', ' ⊂ '),
        (r'\\subseteq', ' ⊆ '),
        (r'\\in', ' ∈ '),
        (r'\\forall', '∀'),
        (r'\\exists', '∃'),
        (r'\\ge\b|\\geq\b', ' ≥ '),
        (r'\\le\b|\\leq\b', ' ≤ '),
        (r'\\ne\b|\\neq\b', ' ≠ '),
        (r'\\succ\b', ' ≻ '),
        (r'\\succeq\b', ' ⪰ '),
        (r'\\approx', ' ≈ '),
        (r'\\sim', ' ~ '),
        (r'\\to', ' → '),
        (r'\\implies', ' ⟹ '),
        (r'\\longrightarrow', ' ⟶ '),
        (r'\\Theta', 'Θ'),
        (r'\\Omega', 'Ω'),
        (r'\\beta', 'β'),
        (r'\\sigma_j\^\{\(c\)\}?', 'σ_j^(c)'),
        (r'\\sigma\^\{\(c\)\}?', 'σ^(c)'),
        (r'\\sigma', 'σ'),
        (r'\\mu_\{?\\text\{norm\}\}?', 'μ_norm'),
        (r'\\mu', 'μ'),
        (r'\\kappa', 'κ'),
        (r'\\rho', 'ρ'),
        (r'\\lambda_\{?\\text\{min\}\}?', 'λ_min'),
        (r'\\lambda_\{?\\text\{max\}\}?', 'λ_max'),
        (r'\\lambda', 'λ'),
        (r'\\mathbb\{R\}\^N', 'ℝ^N'),
        (r'\\mathbb\{R\}', 'ℝ'),
        (r'\\mathbb\{F\}_2', '𝔽₂'),
        (r'\\mathbb\{E\}', '𝔼'),
        (r'\\text\{int\}', 'int'),
        (r'\\text\{relint\}', 'relint'),
        (r'\\text\{rank\}', 'rank'),
        (r'\\text\{diag\}', 'diag'),
        (r'\\text\{Tr\}', 'Tr'),
        (r'\\text\{Var\}', 'Var'),
        (r'\\text\{sign\}', 'sign'),
        (r'\\text\{Vol\}', 'Vol'),
        (r'\\text\{dim\}', 'dim'),
        (r'\\text\{const\}', 'const'),
        (r'\\text\{Faces\}_d', 'Faces_d'),
        (r'\\ker', 'ker'),
        (r'\\liminf', 'lim inf'),
        (r'\\limsup', 'lim sup'),
        (r'\\lim', 'lim'),
        (r'\\left\(', '('), (r'\\right\)', ')'),
        (r'\\left\[', '['), (r'\\right\]', ']'),
        (r'\\left\\\{', '{'), (r'\\right\\\}', '}'),
        (r'\\sum_\{c=1\}\^M', '∑_{c=1}^M'),
        (r'\\sum', '∑'),
        (r'\\prod', '∏'),
        (r'\\dots', '...'),
        (r'\\blacksquare', '■'),
        (r'\\qquad', '    '),
        (r'\\quad', '  '),
        (r'\\times', ' × '),
        (r'\\cdot', ' · '),
        (r'\\partial_k', '∂_k'),
        (r'\\partial_i', '∂_i'),
        (r'\\partial', '∂'),
        (r'\\infty', '∞'),
        (r'\\pm', '±'),
        (r'\\oplus', '⊕'),
        (r'\\max', 'max'),
        (r'\\min', 'min'),
        (r'\\ln', 'ln'),
        (r'\\log', 'log'),
        (r'\\ blabla_placeholder', ''),
        (r'\\,', ' '),
        (r'\\;', ' '),
        (r'\\:', ' '),
        (r'\\!', ''),
    ]
    for p, r in repl:
        text = re.sub(p, r, text)

    # Fractions: \frac{a}{b} -> a/b if simple or (a/b)
    text = re.sub(r'\\frac\{([0-9a-zA-Z]+)\}\{([0-9a-zA-Z]+)\}', r'\1/\2', text)
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1/\2)', text)
    text = re.sub(r'\(\(([^)]+)\)\)', r'(\1)', text)
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def add_formatted_runs(paragraph, text, base_italic=False, base_color=None):
    """Parses inline bold (**...**) and italic (*...*) into formatted runs."""
    # Pattern to match **bold** or *italic*
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = paragraph.add_run(content)
            run.bold = True
            run.italic = base_italic
        elif part.startswith('*') and part.endswith('*'):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.italic = base_italic
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        if base_color:
            run.font.color.rgb = base_color

def parse_line_math(line):
    """Converts inline LaTeX $...$ to cleaned math."""
    def repl_inline(m):
        raw = m.group(1)
        return clean_math(raw)
    return re.sub(r'\$([^$]+)\$', repl_inline, line)

def create_doc():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    title = doc.add_heading(level=1)
    run_t = title.add_run("Resposta Técnica ao Parecer nº 10: Consolidação do Framework CLG-R (Versão 2.0)")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(16, 44, 87)
    
    with open(r"C:\MathDoCarvalho\P_NP\Publicacoes\RespostaAoProfessor_Analise10.md", "r", encoding="utf-8") as f:
        raw_text = f.read()
        
    lines = raw_text.split("\n")
    in_math_block = False
    math_block_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        
        if not line or line.startswith("# Resposta Técnica"):
            continue
        if line.startswith("---"):
            continue
            
        # Check for math display block $$ ... $$
        if line.startswith("$$"):
            if line.endswith("$$") and len(line) > 2:
                # Single line $$...$$
                eq_text = line[2:-2].strip()
                eq_clean = clean_math(eq_text)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(eq_clean)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(20, 50, 100)
                continue
            else:
                # Multi-line $$ start
                eq_lines = []
                while i < len(lines):
                    l = lines[i].strip()
                    i += 1
                    if l.endswith("$$"):
                        if len(l) > 2:
                            eq_lines.append(l[:-2].strip())
                        break
                    eq_lines.append(l)
                eq_clean = clean_math(" ".join(eq_lines))
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(eq_clean)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(20, 50, 100)
                continue
                
        # Headings
        if line.startswith("## "):
            h = doc.add_heading(level=2)
            clean_h = parse_line_math(line[3:])
            clean_h = re.sub(r'\*\*|\*', '', clean_h)
            r = h.add_run(clean_h)
            r.font.name = "Calibri"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 72, 140)
            continue
        elif line.startswith("### "):
            h = doc.add_heading(level=3)
            clean_h = parse_line_math(line[4:])
            clean_h = re.sub(r'\*\*|\*', '', clean_h)
            r = h.add_run(clean_h)
            r.font.name = "Calibri"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(40, 40, 40)
            continue
        elif line.startswith("#### "):
            h = doc.add_heading(level=4)
            clean_h = parse_line_math(line[5:])
            clean_h = re.sub(r'\*\*|\*', '', clean_h)
            r = h.add_run(clean_h)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.bold = True
            continue
            
        # Blockquote (lines starting with >)
        if line.startswith(">"):
            content = line[1:].strip()
            content = parse_line_math(content)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, content, base_italic=True, base_color=RGBColor(50, 50, 50))
            continue
            
        # Lists (numbered or bullet)
        cleaned_line = parse_line_math(line)
        
        # Check numbered list (e.g. "1. ", "2. ")
        match_num = re.match(r'^(\d+)\.\s+(.*)', cleaned_line)
        if match_num:
            prefix = match_num.group(1) + ". "
            rest = match_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r_num = p.add_run(prefix)
            r_num.bold = True
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(11)
            add_formatted_runs(p, rest)
            continue
            
        # Check bullet list (e.g. "- " or "* ")
        if cleaned_line.startswith("- ") or cleaned_line.startswith("* "):
            rest = cleaned_line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = p.add_run("•  ")
            r_bullet.bold = True
            r_bullet.font.name = "Calibri"
            r_bullet.font.size = Pt(11)
            add_formatted_runs(p, rest)
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, cleaned_line)
        
    # Save to repository
    doc.save(r"C:\MathDoCarvalho\P_NP\Publicacoes\RespostaAoProfessor_Analise10.docx")
    print("Salvo com sucesso em: Publicacoes/RespostaAoProfessor_Analise10.docx")
    
    # Save to root MathDoCarvalho
    try:
        doc.save(r"C:\MathDoCarvalho\RespostaAoProfessor_Analise10.docx")
        print("Salvo com sucesso em: C:\\MathDoCarvalho\\RespostaAoProfessor_Analise10.docx")
    except PermissionError:
        print("AVISO: O arquivo C:\\MathDoCarvalho\\RespostaAoProfessor_Analise10.docx esta aberto no Microsoft Word.")
        fallback = r"C:\MathDoCarvalho\RespostaAoProfessor_Analise10_Atualizado.docx"
        doc.save(fallback)
        print(f"Versao atualizada salva com sucesso em: {fallback}")

if __name__ == "__main__":
    create_doc()

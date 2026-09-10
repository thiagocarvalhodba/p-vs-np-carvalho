import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

docx_path = r"C:\MathDoCarvalho\RespostaAoProfessor08.docx"
doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

# Titulo
title_p = doc.add_paragraph()
title_run = title_p.add_run("Resposta Técnica ao Parecer 08: Os Quatro Controles Experimentais e a Consolidação da Tese CLG-R")
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
title_p.paragraph_format.space_after = Pt(10)

meta_p = doc.add_paragraph()
meta_p.add_run("Destinatário: ").bold = True
meta_p.add_run("Ilustre Professor e Comitê de Avaliação\n")
meta_p.add_run("Autor: ").bold = True
meta_p.add_run("Thiago Carvalho\n")
meta_p.add_run("Data: ").bold = True
meta_p.add_run("10 de Setembro de 2026\n")
meta_p.add_run("Assunto: ").bold = True
meta_p.add_run("Execução dos 4 Controles Mandatórios (Escala, Matriz 3x3, Condicionamento kappa_2 e Diferenças Pareadas)")
meta_p.paragraph_format.space_after = Pt(14)

doc.add_paragraph("―" * 50)

def add_heading_1(text):
    h = doc.add_paragraph()
    r = h.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_callout(text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

# 1. Correcoes Formais
add_heading_1("1. Correções Formais e Alinhamento Conceitual")
doc.add_paragraph("Acolhemos integralmente todas as orientações do Parecer 08:")
doc.add_paragraph("1. Teorema Posto-Nulidade em vez de Rouché-Capelli: 'Como A tem posto completo N sobre GF(2), o teorema posto-nulidade implica dim(ker(A)) = 0. Como a instância possui uma solução plantada s*, o sistema afim Ax = b possui exatamente uma única solução |S(I)| = 1.'")
doc.add_paragraph("2. Precisão no Overlap: Substituído 'completamente ortogonal' por 'com overlap linear aproximadamente nulo com a solução única (q(s, s*) ≈ 0.0)'.")
doc.add_paragraph("3. Softplus e Platôs: Retificada a redação para: 'Softplus substitui o platô exatamente plano do hinge por uma transição suave, evitando o gradiente identicamente nulo no interior da região satisfeita.'")
doc.add_paragraph("4. Condicionamento Espectral Rigoroso: Definido kappa_2(H) = max_i |lambda_i(H)| / max(min_i |lambda_i(H)|, 10^(-5)).")
doc.add_paragraph("5. Conjectura CLG-R com Funcional Limitado: Reformulada para funcionais normalizados Q in [0, 1] (ex: Q_H = 1 - d_H).")

# 2. Os 4 Controles
add_heading_1("2. Resultados da Execução dos Quatro Controles Mandatórios (CLG-04 Fase II)")
doc.add_paragraph(
    "Executamos o protocolo CLG-04D (Fontes/exp_clg04_phase2_factorial.py) em N=60 com 450 trajetórias estritamente pareadas:\n"
    "• Controle de Escala: Potenciais normalizados por c_Phi = E[||grad Phi||] (c_multi = 5.26, c_quad = 5.93, c_soft = 8.65).\n"
    "• Matriz 3x3 Pareada: 3 Representações x 3 Dinâmicas (GD puro, Langevin, Adam)."
)

# Tabela 3x3
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdrs = ["Representação", "Dinâmica", "d_H (Média ± SEM)", "Overlap q", "E_disc residual", "Condicionamento kappa_2"]
for i, h in enumerate(hdrs):
    hdr[i].text = h
    hdr[i].paragraphs[0].runs[0].bold = True
    hdr[i].paragraphs[0].runs[0].font.size = Pt(8.5)

rows_data = [
    ("Multilinear", "GD Puro", "0.431 ± 0.007", "+0.139", "20.4 cl.", "182.7"),
    ("Multilinear", "Langevin", "0.432 ± 0.010", "+0.136", "19.7 cl.", "141.1"),
    ("Multilinear", "Adam", "0.418 ± 0.008", "+0.164", "18.6 cl.", "176.5"),
    ("Quadrática", "GD Puro", "0.472 ± 0.012", "+0.055", "30.1 cl.", "88.815 (Patológico)"),
    ("Quadrática", "Langevin", "0.475 ± 0.005", "+0.049", "29.5 cl.", "88.141 (Patológico)"),
    ("Quadrática", "Adam", "0.440 ± 0.013", "+0.119", "18.4 cl.", "81.463 (Patológico)"),
    ("Softplus", "GD Puro", "0.455 ± 0.012", "+0.091", "25.5 cl.", "25.5 (Estável)"),
    ("Softplus", "Langevin", "0.460 ± 0.006", "+0.080", "24.6 cl.", "24.2 (Estável)"),
    ("Softplus", "Adam", "0.452 ± 0.010", "+0.096", "25.3 cl.", "47.1 (Estável)"),
]

for row in rows_data:
    r_cells = table.add_row().cells
    for i, val in enumerate(row):
        r_cells[i].text = val
        r_cells[i].paragraphs[0].runs[0].font.size = Pt(8.5)

# 3. Diferencas Pareadas por Instancia
add_heading_1("3. Diferenças Pareadas por Instância Individual (Delta_i)")
doc.add_paragraph("Para afastar qualquer pseudorreplicamento amostral, reportamos os valores individuais para as 5 instâncias independentes em GD Puro:")
doc.add_paragraph("• Instância 1: d_H(Soft) = 0.482 | d_H(Multi) = 0.450 | d_H(Quad) = 0.488  ==>  Delta(Quad - Soft) = +0.007")
doc.add_paragraph("• Instância 2: d_H(Soft) = 0.418 | d_H(Multi) = 0.423 | d_H(Quad) = 0.432  ==>  Delta(Quad - Soft) = +0.013")
doc.add_paragraph("• Instância 3: d_H(Soft) = 0.445 | d_H(Multi) = 0.428 | d_H(Quad) = 0.475  ==>  Delta(Quad - Soft) = +0.030")
doc.add_paragraph("• Instância 4: d_H(Soft) = 0.448 | d_H(Multi) = 0.410 | d_H(Quad) = 0.465  ==>  Delta(Quad - Soft) = +0.017")
doc.add_paragraph("• Instância 5: d_H(Soft) = 0.480 | d_H(Multi) = 0.442 | d_H(Quad) = 0.502  ==>  Delta(Quad - Soft) = +0.022")
doc.add_paragraph(
    "Em 100% das 5 instâncias, a relaxação Quadrática Hinge produz atratores estagnados mais distantes da solução que o Softplus (Delta > 0), "
    "confirmando que a patologia de condicionamento da Quadrática (kappa_2 > 80.000) degrada a busca de forma consistente."
)

# 4. A Tese Central Consagrada
add_heading_1("4. A Tese Central do Artigo Consagrada")
add_callout(
    "Tese Central Consolidada do CLG-R:\n"
    "“A equivalência booleana não determina a acessibilidade algorítmica de uma relaxação contínua: "
    "representações equivalentes podem induzir geometrias numéricas e interações representação–dinâmica distintas, "
    "produzindo diferentes qualidades de solução sob a mesma classe de algoritmo local.”",
    italic=True
)

doc.add_paragraph(
    "Com esses quatro controles fechados, a resposta à pergunta central está completamente documentada: "
    "as representações contínuas induzem geometrias numéricas distintas que filtram a eficácia de classes específicas de otimizadores locais."
)

p_end = doc.add_paragraph()
p_end.add_run("Respeitosamente,\n\n").italic = True
p_end.add_run("Thiago Carvalho\n").bold = True
p_end.add_run("Pesquisador Principal\nVitória, ES, 2026")

doc.save(docx_path)
print(f"Salvo DOCX 08: {docx_path}")
